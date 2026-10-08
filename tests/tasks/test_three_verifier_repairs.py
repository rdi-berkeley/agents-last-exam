import math
from copy import deepcopy
from io import BytesIO
from xml.etree import ElementTree as ET

import pytest
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from tasks.business_finance.equity_research_summary.scripts.score_workbook import (
    score_workbook_bytes,
)
from tasks.life_sciences.amber_three_stage_mmgbsa_workflow_instance_1.scripts.verify_submission import (
    evaluate_output_bundle,
)
from tasks.life_sciences.tp53_locus_variant_histone_browser_svg.scripts.score_svg import (
    score_svg_bytes,
)


@pytest.fixture
def signal_svg():
    signal = [50 + 35 * math.sin(i * 0.7) + 20 * math.cos(i * 0.21) for i in range(60)]
    reference = {"chrom": "chr17", "start": 7571651, "end": 7590910, "signal_profile": signal}
    points = " ".join(f"{i * 10},{200 - v}" for i, v in enumerate(signal))
    svg = f'<svg><text>chr17 7571651 7580000 7590910 hg19 VCF H3K27ac</text><polygon points="{points}"/></svg>'
    return svg, reference


def test_polygon_matches_same_profile_as_polyline(signal_svg):
    svg, ref = signal_svg
    assert score_svg_bytes(svg.encode(), ref).score == 1
    assert score_svg_bytes(svg.replace("polygon", "polyline").encode(), ref).score == 1


@pytest.mark.parametrize(
    "defect",
    [
        "flat",
        "wrong_signal",
        "few_points",
        "odd_points",
        "nan",
        "overflow",
        "hidden",
        "transparent",
        "none",
        "defs",
        "transform",
        "wrong_locus",
        "no_tracks",
        "too_many",
        "style_none",
        "inherited_none",
        "zero_exponent",
        "stylesheet",
        "nested_svg",
        "animation",
    ],
)
def test_invalid_polygon_does_not_gain_full_score(signal_svg, defect):
    svg, ref = signal_svg
    root = ET.fromstring(svg)
    polygon = root[1]
    if defect == "flat":
        polygon.set("points", " ".join(f"{i * 10},100" for i in range(60)))
    elif defect == "wrong_signal":
        polygon.set(
            "points", " ".join(f"{i * 10},{100 + 50 * math.sin(i * 1.7)}" for i in range(60))
        )
    elif defect == "few_points":
        polygon.set("points", "0,0 100,100 590,20")
    elif defect in {"odd_points", "nan", "overflow"}:
        polygon.set(
            "points",
            polygon.get("points")
            + {"odd_points": " 1", "nan": " nan,20", "overflow": " 1e999,2"}[defect],
        )
    elif defect in {"hidden", "transparent", "none", "transform"}:
        attr, value = {
            "hidden": ("display", "none"),
            "transparent": ("opacity", "0"),
            "none": ("fill", "none"),
            "transform": ("transform", "scale(0)"),
        }[defect]
        polygon.set(attr, value)
    elif defect == "defs":
        root.remove(polygon)
        ET.SubElement(root, "defs").append(polygon)
    elif defect == "style_none":
        polygon.set("style", "fill: none")
    elif defect == "inherited_none":
        root.set("fill", "none")
    elif defect == "zero_exponent":
        root.set("opacity", "0e10")
    elif defect == "stylesheet":
        ET.SubElement(root, "style").text = "polygon { display:none; }"
    elif defect == "nested_svg":
        root.remove(polygon)
        ET.SubElement(root, "svg", width="0").append(polygon)
    elif defect == "animation":
        ET.SubElement(polygon, "set", attributeName="display", to="none")
    elif defect == "wrong_locus":
        root[0].text = "chr17 1 2 3 hg19 VCF H3K27ac"
    elif defect == "no_tracks":
        root[0].text = "chr17 7571651 7580000 7590910 hg19"
    else:
        for _ in range(20):
            root.append(deepcopy(polygon))
    assert score_svg_bytes(ET.tostring(root), ref).score < 1


@pytest.fixture
def workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A2"] = "Market Data: source note"
    ws["A4"] = "Market Data (from Yahoo Finance)"
    ws["A4"].font = Font(bold=True)
    ws["A4"].fill = PatternFill("solid", fgColor="123456")
    ws["A5"], ws["B5"] = "Revenue", 100
    ws["A6"], ws["B6"] = "Margin", "=B5/100"
    manifest = {
        "sheet_name_contains": "Summary",
        "section_headers": [{"text": "Market Data", "bold": True, "fill": True}],
        "fixed_values": [{"label": "Revenue", "value": 100, "tolerance": 1}],
        "formula_cells": [{"label": "Margin", "must_reference": ["Revenue"]}],
    }
    return wb, manifest


@pytest.mark.parametrize(
    "defect",
    [
        None,
        "no_bold",
        "no_fill",
        "split_styles",
        "missing_title",
        "prose_title",
        "wrong_number",
        "wrong_formula",
    ],
)
def test_workbook_requires_one_real_fully_formatted_header(workbook, defect):
    wb, manifest = workbook
    ws = wb.active
    if defect in {"no_bold", "split_styles"}:
        ws["A4"].font = Font(bold=False)
    if defect == "no_fill":
        ws["A4"].fill = PatternFill()
    if defect == "split_styles":
        ws["A7"] = "Market Data"
        ws["A7"].font = Font(bold=True)
    if defect == "missing_title":
        ws["A4"] = "Absent"
    if defect == "prose_title":
        ws["A4"] = "No Market Data was supplied"
    if defect == "wrong_number":
        ws["B5"] = 102
    if defect == "wrong_formula":
        ws["B6"] = "=B7/100"
    stream = BytesIO()
    wb.save(stream)
    assert (score_workbook_bytes(stream.getvalue(), manifest).score == 1) == (defect is None)


@pytest.fixture
def amber_bundle():
    basename = "GLN_phb2_parl_pgam5_model_0"
    common = f'''#!/bin/bash
#SBATCH --nodes=1
NAME="{basename}"
C="/TASK/params/${{NAME}}.prmtop"
R="/TASK/params/${{NAME}}_receptor.prmtop"
L="/TASK/params/${{NAME}}_ligand.prmtop"
'''
    build = """if [[ ! -f "$C" ]]; then
cat > /TASK/build.leap <<EOF
source leaprc.protein.ff14SB
set default PBRadii mbondi3
c = loadpdb /TASK/input/complex_structure.pdb
saveamberparm c ${C} /TASK/params/complex.inpcrd
quit
EOF
tleap -f /TASK/build.leap
fi
if [[ ! -f "$R" || ! -f "$L" ]]; then
ante-MMPBSA.py -p "$C" -r "$R" -l "$L" -n :3-7
fi
pmemd.cuda -O -o min.out -r min.rst
pmemd.cuda -O -o equil.out -r equil.rst
"""
    mm = """cat > /TASK/mmgbsa.in <<EOF
&general
startframe=1, endframe=250, interval=1,
/
&gb
igb=8, saltcon=0.150,
/
EOF
MMPBSA.py -O -i /TASK/mmgbsa.in -o /TASK/output/FINAL_RESULTS_MMGBSA.dat -cp "$C" -rp "$R" -lp "$L" -y /TASK/input/prod.mdcrd
"""
    files = {
        "submit_min.sh": common + build,
        "submit_prod.sh": common
        + "pmemd.cuda -O -c equil.rst -o prod.out -r prod.rst -x prod.mdcrd\n",
        "submit_mmgbsa.sh": common + mm,
        "FINAL_RESULTS_MMGBSA.dat": "Calculations performed using 250.0 complex frames.\nReceptor mask :1-2\nLigand mask :3-7\nDELTA TOTAL -116.3\n",
    }
    pdb = "\n".join(
        f"ATOM  {i:5d}  CA  ALA {chain}{i:4d}    1.000 2.000 3.000"
        for i, chain in enumerate("AABBBCC", 1)
    )
    return files, pdb


@pytest.mark.parametrize(
    "defect",
    [
        None,
        "wrong_mask",
        "wrong_role",
        "duplicate_option",
        "late_value",
        "single_quote",
        "comment_split",
        "echo_split",
        "dead_branch",
        "function",
        "after_exit",
        "wrong_structure",
        "wrong_trajectory",
        "wrong_energy",
        "missing_result",
        "missing_prod",
        "wrong_igb",
        "few_frames",
        "strip_extra",
        "top_overwrite",
        "top_delete",
        "leap_remove",
        "prod_comment",
        "prod_bad_restart",
        "prod_overwrite",
        "dry_shell",
        "compound_remove",
        "active_fallback",
        "wrong_result_mask",
        "one_frame_result",
        "echo_substitution",
        "conditional_halt",
        "extra_deliverable",
        "unterminated_quote",
        "unterminated_heredoc",
    ],
)
def test_cross_stage_split_preserves_all_scientific_and_artifact_gates(amber_bundle, defect):
    files, pdb = amber_bundle
    if defect == "wrong_mask":
        files["submit_min.sh"] = files["submit_min.sh"].replace("-n :3-7", "-n :1-2")
    elif defect in {"comment_split", "echo_split"}:
        files["submit_min.sh"] = files["submit_min.sh"].replace(
            "ante-MMPBSA.py -p",
            ("# " if defect == "comment_split" else "echo ") + "ante-MMPBSA.py -p",
        )
    elif defect == "dead_branch":
        files["submit_min.sh"] = files["submit_min.sh"].replace(
            'if [[ ! -f "$R" || ! -f "$L" ]]; then', "if false; then"
        )
    elif defect == "function":
        files["submit_min.sh"] = (
            files["submit_min.sh"]
            .replace("ante-MMPBSA.py -p", "never_called() { ante-MMPBSA.py -p")
            .replace("-n :3-7", "-n :3-7; }")
        )
    elif defect == "after_exit":
        files["submit_mmgbsa.sh"] = files["submit_mmgbsa.sh"].replace(
            "MMPBSA.py -O", "exit 0\nMMPBSA.py -O"
        )
    elif defect in {"wrong_role", "duplicate_option", "single_quote"}:
        replacement = {
            "wrong_role": '-rp "$L"',
            "duplicate_option": '-rp "$R" -rp "$R"',
            "single_quote": "-rp '$R'",
        }[defect]
        files["submit_mmgbsa.sh"] = files["submit_mmgbsa.sh"].replace('-rp "$R"', replacement)
    elif defect == "late_value":
        files["submit_mmgbsa.sh"] = (
            files["submit_mmgbsa.sh"].replace('-rp "$R"', '-rp "$LATE"') + '\nLATE="$R"\n'
        )
    elif defect == "wrong_structure":
        files["submit_min.sh"] = files["submit_min.sh"].replace(
            "input/complex_structure.pdb", "input/another.pdb"
        )
    elif defect == "wrong_trajectory":
        files["submit_mmgbsa.sh"] = files["submit_mmgbsa.sh"].replace(
            "/TASK/input/prod.mdcrd", "/TASK/prod.mdcrd"
        )
    elif defect == "wrong_energy":
        files["FINAL_RESULTS_MMGBSA.dat"] = files["FINAL_RESULTS_MMGBSA.dat"].replace(
            "-116.3", "-200"
        )
    elif defect in {"missing_result", "missing_prod"}:
        del files["FINAL_RESULTS_MMGBSA.dat" if defect == "missing_result" else "submit_prod.sh"]
    elif defect in {"wrong_igb", "few_frames"}:
        old, new = ("igb=8", "igb=7") if defect == "wrong_igb" else ("endframe=250", "endframe=2")
        files["submit_mmgbsa.sh"] = files["submit_mmgbsa.sh"].replace(old, new)
    elif defect == "strip_extra":
        files["submit_min.sh"] = files["submit_min.sh"].replace("-n :3-7", "-n :3-7 -s :1")
    elif defect == "top_overwrite":
        files["submit_min.sh"] += '\necho broken > "$R"\n'
    elif defect == "top_delete":
        files["submit_min.sh"] += '\nif [[ ! -f "$C" ]]; then\nrm -f "$R"\nfi\n'
    elif defect == "leap_remove":
        files["submit_min.sh"] = files["submit_min.sh"].replace(
            "saveamberparm c", "remove c c.1\nsaveamberparm c"
        )
    elif defect == "prod_comment":
        files["submit_prod.sh"] = files["submit_prod.sh"].replace("pmemd.cuda", "# pmemd.cuda")
    elif defect == "prod_bad_restart":
        files["submit_prod.sh"] = files["submit_prod.sh"].replace(
            "-c equil.rst", "-c min_nonexistent.rst"
        )
    elif defect == "prod_overwrite":
        files["submit_prod.sh"] += '\necho invalid > "$R"\n'
    elif defect == "dry_shell":
        files["submit_min.sh"] = files["submit_min.sh"].replace("NAME=", "set -n\nNAME=")
    elif defect == "compound_remove":
        files["submit_min.sh"] += '\necho done; rm -f "$R"\n'
    elif defect == "active_fallback":
        files["submit_mmgbsa.sh"] = (
            files["submit_mmgbsa.sh"]
            .replace(
                "MMPBSA.py -O",
                'T=/TASK/input/prod.mdcrd\nif [[ -s "$T" ]]; then\nT=/TASK/wrong.mdcrd\nfi\nMMPBSA.py -O',
            )
            .replace("-y /TASK/input/prod.mdcrd", '-y "$T"')
        )
    elif defect == "wrong_result_mask":
        files["FINAL_RESULTS_MMGBSA.dat"] = files["FINAL_RESULTS_MMGBSA.dat"].replace(
            "Ligand mask :3-7", "Ligand mask :1-2"
        )
    elif defect == "one_frame_result":
        files["FINAL_RESULTS_MMGBSA.dat"] = files["FINAL_RESULTS_MMGBSA.dat"].replace(
            "250.0 complex", "1.0 complex"
        )
    elif defect == "echo_substitution":
        files["submit_min.sh"] += '\necho "$(rm -f $R)"\n'
    elif defect == "conditional_halt":
        files["submit_min.sh"] = files["submit_min.sh"].replace(
            "NAME=", "if true; then\nexit 1\nfi\nNAME="
        )
    elif defect == "extra_deliverable":
        files["unrequested.txt"] = "extra file"
    elif defect == "unterminated_quote":
        files["submit_min.sh"] += '\necho "unfinished\n'
    elif defect == "unterminated_heredoc":
        files["submit_min.sh"] += "\ncat > /TASK/extra.in <<EOF\nunfinished\n"
    result = evaluate_output_bundle(files, input_pdb_text=pdb, task_dir="/TASK")
    assert (result["score"] == 1) == (defect is None), result


@pytest.mark.parametrize(
    "layout", ["absolute", "relative", "slurm", "script_dir", "alternate_dirs"]
)
def test_cross_stage_uses_real_task_metadata(amber_bundle, layout):
    files, pdb = amber_bundle
    task_dir = (
        "/media/user/data/agenthle/life_sciences/amber_three_stage_mmgbsa_workflow_instance_1/base"
    )
    input_dir, output_dir = task_dir + "/input", task_dir + "/output"
    if layout == "alternate_dirs":
        input_dir, output_dir = "/mounted/inputs", "/mounted/deliverables"
    for name, text in files.items():
        text = text.replace("/TASK/input", input_dir).replace("/TASK/output", output_dir)
        text = text.replace("/TASK", task_dir)
        if layout == "relative":
            text = text.replace(task_dir + "/", "./")
        elif layout == "slurm":
            text = text.replace(task_dir, "${SLURM_SUBMIT_DIR}")
        elif layout == "script_dir" and name.endswith(".sh"):
            text = text.replace(task_dir, "${ROOT}")
            text = text.replace(
                "#SBATCH --nodes=1\n",
                "#SBATCH --nodes=1\n"
                'DIR="$(cd "$(dirname "${BASH_SOURCE[0]}")" && pwd)"\n'
                'ROOT="$(cd "${DIR}/.." && pwd)"\n',
            )
        files[name] = text
    result = evaluate_output_bundle(
        files, input_pdb_text=pdb, task_dir=task_dir, input_dir=input_dir, output_dir=output_dir
    )
    assert result["score"] == 1, result


@pytest.mark.parametrize(
    "path",
    [
        "/TASK/input/prod.mdcrd",
        "/real-task-other/input/prod.mdcrd",
        "/real-task/input/../../wrong/prod.mdcrd",
        "/real-task/output/prod.mdcrd",
    ],
)
def test_cross_stage_rejects_other_trajectory_roots(amber_bundle, path):
    files, pdb = amber_bundle
    files = {name: text.replace("/TASK", "/real-task") for name, text in files.items()}
    files["submit_mmgbsa.sh"] = files["submit_mmgbsa.sh"].replace(
        "/real-task/input/prod.mdcrd", path
    )
    result = evaluate_output_bundle(files, input_pdb_text=pdb, task_dir="/real-task")
    assert result["score"] == 0
    assert "provided input trajectory" in result["cross_stage_split"]["reason"]


def test_cross_stage_does_not_infer_paths_from_candidate(amber_bundle):
    files, pdb = amber_bundle
    result = evaluate_output_bundle(files, input_pdb_text=pdb)
    assert result["score"] == 0
    assert "task directory" in result["cross_stage_split"]["reason"]
