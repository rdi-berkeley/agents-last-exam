from io import BytesIO
from zipfile import ZipFile

import pytest
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_from_string

from tasks.business_finance.equity_research_summary.scripts.formula_evaluator import (
    FormulaEvaluationError,
    FormulaEvaluator,
)
from tasks.business_finance.equity_research_summary.scripts.score_workbook import (
    score_workbook_bytes,
)

INPUTS = {
    "revenue": ("Total Revenue", 100),
    "cost": ("Cost of Revenue", 60),
    "cash_flow": ("Operating Cash Flow", 40),
    "capex": ("Capital Expenditures", 10),
    "market_cap": ("Market Cap ($B)", 2),
}
OUTPUTS = {
    "gross_profit": ("Gross Profit", "={revenue}-{cost}"),
    "free_cash_flow": ("Free Cash Flow", "={cash_flow}-{capex}"),
    "market_cap_revenue": ("Market Cap / Revenue (x)", "={market_cap}*1000/{revenue}"),
    "cash_flow_yield": ("Free Cash Flow Yield", "={free_cash_flow}/({market_cap}*1000)"),
}
LAYOUTS = {
    "vertical": {key: (row, 1) for row, key in enumerate([*INPUTS, *OUTPUTS], start=2)},
    "rearranged": {
        "revenue": (8, 4),
        "cost": (3, 7),
        "cash_flow": (11, 2),
        "capex": (2, 10),
        "market_cap": (5, 4),
        "gross_profit": (4, 7),
        "free_cash_flow": (9, 10),
        "market_cap_revenue": (6, 2),
        "cash_flow_yield": (12, 7),
    },
}


@pytest.fixture(params=LAYOUTS, ids=LAYOUTS)
def workbook_case(request):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "TSLA Summary"
    addresses = {}
    for key, (row, column) in LAYOUTS[request.param].items():
        label, _ = (INPUTS | OUTPUTS)[key]
        sheet.cell(row, column, label)
        addresses[key] = f"{get_column_letter(column + 1)}{row}"
    for key, (_, value) in INPUTS.items():
        sheet[addresses[key]] = value
    for key, (_, formula) in OUTPUTS.items():
        sheet[addresses[key]] = formula.format(**addresses)
    manifest = {
        "sheet_name_contains": "TSLA",
        "fixed_values": [
            {"label": label, "value": value, "tolerance": 0}
            for key, (label, value) in INPUTS.items()
            if key != "market_cap"
        ],
        "live_labels": ["Market Cap ($B)"],
        "formula_cells": [
            {"label": "Gross Profit", "must_reference": ["Total Revenue", "Cost of Revenue"]},
            {
                "label": "Free Cash Flow",
                "must_reference": ["Operating Cash Flow", "Capital Expenditures"],
            },
        ],
        "dashboard_formulas": [
            {
                "label": "Market Cap / Revenue",
                "must_reference": ["Market Cap ($B)", "Total Revenue"],
            },
            {
                "label": "Free Cash Flow Yield",
                "must_reference": ["Free Cash Flow", "Market Cap ($B)"],
            },
        ],
    }
    return workbook, manifest, addresses


def score_case(workbook_case):
    workbook, manifest, _ = workbook_case
    buffer = BytesIO()
    workbook.save(buffer)
    return score_workbook_bytes(buffer.getvalue(), manifest)


def test_canonical_formulas_pass_in_each_layout(workbook_case):
    result = score_case(workbook_case)
    assert result.score == 1.0
    assert result.passed
    assert not result.reasons


@pytest.mark.parametrize(
    "formula",
    [
        "=({cash_flow})+(-{capex})",
        "=-({capex}-{cash_flow})",
        "={cash_flow_absolute}-{capex_absolute}",
        "=SUM({cash_flow},-{capex})",
        "=IF(TRUE,{cash_flow}-{capex},0)",
        "=IF(AND(TRUE(),NOT(FALSE())),{cash_flow}-{capex},0)",
    ],
    ids=[
        "addition-of-negative",
        "reverse-subtraction",
        "absolute",
        "sum",
        "if",
        "boolean-functions",
    ],
)
def test_equivalent_expressions_and_functions_pass(workbook_case, formula):
    workbook, _, addresses = workbook_case
    references = dict(addresses)
    for key in ("cash_flow", "capex"):
        column, row = coordinate_from_string(addresses[key])
        references[f"{key}_absolute"] = f"${column}${row}"
    workbook.active[addresses["free_cash_flow"]] = formula.format(**references)
    result = score_case(workbook_case)
    assert result.score == 1.0
    assert result.passed


@pytest.mark.parametrize(
    "output,formula",
    [
        ("market_cap_revenue", "={market_cap}/({revenue}/1000)"),
        ("cash_flow_yield", "={free_cash_flow}/{market_cap}/1000"),
    ],
)
def test_equivalent_unit_conversions_pass(workbook_case, output, formula):
    workbook, _, addresses = workbook_case
    workbook.active[addresses[output]] = formula.format(**addresses)
    result = score_case(workbook_case)
    assert result.score == 1.0
    assert result.passed


@pytest.mark.parametrize("formula", ["=X1+X2", '=INDIRECT("X1")+INDIRECT("X2")'])
def test_references_through_helper_cells_pass(workbook_case, formula):
    workbook, _, addresses = workbook_case
    sheet = workbook.active
    sheet["X1"] = f"={addresses['cash_flow']}"
    sheet["X2"] = f"=-{addresses['capex']}"
    sheet[addresses["free_cash_flow"]] = formula
    result = score_case(workbook_case)
    assert result.score == 1.0
    assert result.passed


@pytest.mark.parametrize(
    "formula",
    ["={cash_flow}+{capex}", "={capex}-{cash_flow}"],
    ids=["addition", "reversed-subtraction"],
)
def test_wrong_cash_flow_arithmetic_is_not_full_credit(workbook_case, formula):
    workbook, _, addresses = workbook_case
    workbook.active[addresses["free_cash_flow"]] = formula.format(**addresses)
    result = score_case(workbook_case)
    assert result.score < 1.0
    assert not result.passed


@pytest.mark.parametrize(
    "output,formula",
    [
        ("market_cap_revenue", "={market_cap}/{revenue}"),
        ("cash_flow_yield", "={free_cash_flow}/{market_cap}"),
    ],
    ids=["market-cap-revenue", "cash-flow-yield"],
)
def test_missing_billions_to_millions_conversion_is_not_full_credit(workbook_case, output, formula):
    workbook, _, addresses = workbook_case
    workbook.active[addresses[output]] = formula.format(**addresses)
    result = score_case(workbook_case)
    assert result.score < 1.0
    assert not result.passed


def test_cell_addresses_inside_strings_do_not_count_as_references(workbook_case):
    workbook, _, addresses = workbook_case
    workbook.active[addresses["free_cash_flow"]] = '=30+0*LEN("{cash_flow} {capex}")'.format(
        **addresses
    )
    result = score_case(workbook_case)
    assert result.score < 1.0
    assert not result.passed


@pytest.mark.parametrize(
    "formula",
    [
        "=40-{capex}",
        "={cash_flow}-10",
        "=40-{capex}+0*{cash_flow}",
        "={cash_flow}-10+0*{capex}",
        "=30+0*{cash_flow}+0*{capex}",
    ],
    ids=[
        "cash-flow-literal",
        "capex-literal",
        "cash-flow-zero-weight",
        "capex-zero-weight",
        "result-zero-weight",
    ],
)
def test_hardcoded_inputs_are_not_full_credit(workbook_case, formula):
    workbook, _, addresses = workbook_case
    workbook.active[addresses["free_cash_flow"]] = formula.format(**addresses)
    result = score_case(workbook_case)
    assert result.score < 1.0
    assert not result.passed


def test_hardcoded_result_is_not_full_credit(workbook_case):
    workbook, _, addresses = workbook_case
    workbook.active[addresses["free_cash_flow"]] = 30
    result = score_case(workbook_case)
    assert result.score < 1.0
    assert not result.passed


def test_formula_wrong_beyond_first_perturbation_is_not_full_credit(workbook_case):
    workbook, _, addresses = workbook_case
    cash_flow = addresses["cash_flow"]
    capex = addresses["capex"]
    workbook.active[addresses["free_cash_flow"]] = (
        f"={cash_flow}-{capex}+({cash_flow}-40)*({cash_flow}-59.18)"
    )
    assert score_case(workbook_case).score < 1.0


def test_formula_wrong_only_when_inputs_change_together_is_not_full_credit(workbook_case):
    workbook, _, addresses = workbook_case
    cash_flow = addresses["cash_flow"]
    capex = addresses["capex"]
    workbook.active[addresses["free_cash_flow"]] = (
        f"={cash_flow}-{capex}+({cash_flow}-40)*({capex}-10)"
    )
    assert score_case(workbook_case).score < 1.0


def test_formula_cannot_target_known_perturbations(workbook_case):
    workbook, _, addresses = workbook_case
    cash_flow = addresses["cash_flow"]
    capex = addresses["capex"]
    workbook.active[addresses["free_cash_flow"]] = (
        f"={cash_flow}-{capex}+({cash_flow}-40)*({capex}-10)"
        f"*({cash_flow}-59.18)*({cash_flow}-41.82)"
    )
    assert score_case(workbook_case).score < 1.0


def test_rounding_cannot_hide_a_hardcoded_formula(workbook_case):
    workbook, _, addresses = workbook_case
    workbook.active[addresses["free_cash_flow"]] = "=ROUND(30,-10)"
    assert score_case(workbook_case).score < 1.0


def test_coarse_rounding_cannot_hide_the_result(workbook_case):
    workbook, _, addresses = workbook_case
    workbook.active[addresses["free_cash_flow"]] = (
        f"=ROUND({addresses['cash_flow']}-{addresses['capex']},-10)"
    )
    result = score_case(workbook_case)
    assert "formula_value_mismatch:Free Cash Flow" in result.reasons


def test_nonzero_rounded_result_is_accepted(workbook_case):
    workbook, manifest, addresses = workbook_case
    workbook.active[addresses["cash_flow"]] = 41
    next(entry for entry in manifest["fixed_values"] if entry["label"] == "Operating Cash Flow")[
        "value"
    ] = 41
    workbook.active[addresses["free_cash_flow"]] = (
        f"=ROUND({addresses['cash_flow']}-{addresses['capex']},-1)"
    )
    assert score_case(workbook_case).score == 1.0


def test_iferror_handles_an_error_cell(workbook_case):
    workbook, _, addresses = workbook_case
    sheet = workbook.active
    sheet["X1"] = "#DIV/0!"
    sheet[addresses["free_cash_flow"]] = (
        f"=IFERROR(X1,{addresses['cash_flow']}-{addresses['capex']})"
    )
    assert score_case(workbook_case).score == 1.0


@pytest.mark.parametrize("precision", ["2", "2.0"])
@pytest.mark.parametrize("space", ["", " "])
def test_rounding_accepts_equivalent_precision_and_whitespace(workbook_case, precision, space):
    workbook, _, addresses = workbook_case
    workbook.active[addresses["free_cash_flow"]] = (
        f"=ROUND({addresses['cash_flow']}-{addresses['capex']}{space},{precision})"
    )
    assert score_case(workbook_case).score == 1.0


def test_iferror_can_handle_error_in_helper_cell(workbook_case):
    workbook, _, addresses = workbook_case
    sheet = workbook.active
    sheet["X1"] = "=1/0"
    sheet[addresses["free_cash_flow"]] = (
        f"=IFERROR(X1,{addresses['cash_flow']}-{addresses['capex']})"
    )
    assert score_case(workbook_case).score == 1.0


def test_invalid_fixed_value_yields_partial_credit_not_evaluator_error(workbook_case):
    workbook, _, addresses = workbook_case
    workbook.active[addresses["cash_flow"]] = "N/A"
    result = score_case(workbook_case)
    assert "fixed_value_not_numeric:Operating Cash Flow" in result.reasons
    assert "formula_calculation_error:Free Cash Flow" in result.reasons
    assert result.score < 1.0


def test_formula_range_does_not_materialize_empty_cells():
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "=SUM(Z1:Z10000)"
    evaluator = FormulaEvaluator(workbook, sheet)
    assert evaluator.numeric(sheet["A1"], {}, {}) == 0
    assert set(sheet._cells) == {(1, 1)}


def test_many_ranges_cannot_bypass_reference_budget():
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "=SUM(Z1:Z10000)+SUM(AA1:AA10000)+SUM(AB1:AB10000)"
    evaluator = FormulaEvaluator(workbook, sheet)
    with pytest.raises(FormulaEvaluationError, match="reference is too large"):
        evaluator.numeric(sheet["A1"], {}, {})
    assert set(sheet._cells) == {(1, 1)}


def test_unsupported_legal_function_raises_instead_of_scoring_zero(workbook_case):
    workbook, _, addresses = workbook_case
    sheet = workbook.active
    assert score_case(workbook_case).score == 1.0
    sheet[addresses["free_cash_flow"]] = (
        "=_xlfn.LET(_xlpm.flow,{cash_flow},_xlpm.cost,{capex},_xlpm.flow-_xlpm.cost)".format(
            **addresses
        )
    )
    try:
        result = score_case(workbook_case)
    except FormulaEvaluationError:
        return

    assert result.score == 1.0
    assert result.passed
    sheet[addresses["free_cash_flow"]] = (
        "=_xlfn.LET(_xlpm.flow,{cash_flow},_xlpm.cost,{capex},_xlpm.flow+_xlpm.cost)".format(
            **addresses
        )
    )
    assert score_case(workbook_case).score < 1.0


@pytest.mark.parametrize("function", ["ROUND", "ROUNDUP", "ROUNDDOWN", "TRUNC"])
def test_explicit_output_rounding_is_accepted(workbook_case, function):
    workbook, _, addresses = workbook_case
    workbook.active[addresses["market_cap_revenue"]] = (
        f"={function}({addresses['market_cap']}*1000/{addresses['revenue']},2)"
    )
    assert score_case(workbook_case).score == 1.0


def test_wrong_operation_inside_output_rounding_loses_credit(workbook_case):
    workbook, _, addresses = workbook_case
    workbook.active[addresses["free_cash_flow"]] = (
        f"=ROUND({addresses['cash_flow']}+{addresses['capex']},0)"
    )
    assert score_case(workbook_case).score < 1.0


def test_nested_rounding_is_not_incorrectly_scored(workbook_case):
    workbook, _, addresses = workbook_case
    workbook.active[addresses["free_cash_flow"]] = (
        f"=ROUND({addresses['cash_flow']},0)-ROUND({addresses['capex']},0)"
    )
    with pytest.raises(FormulaEvaluationError, match="rounded"):
        score_case(workbook_case)


def test_unsupported_dynamic_reference_is_unscored(workbook_case):
    workbook, _, addresses = workbook_case
    sheet = workbook.active
    sheet["X1"] = addresses["cash_flow"]
    sheet[addresses["free_cash_flow"]] = f"=INDIRECT(X1)-{addresses['capex']}"
    with pytest.raises(FormulaEvaluationError):
        score_case(workbook_case)


def test_formula_that_divides_by_zero_loses_credit(workbook_case):
    workbook, _, addresses = workbook_case
    workbook.active[addresses["free_cash_flow"]] = f"={addresses['cash_flow']}/0"
    result = score_case(workbook_case)
    assert "formula_calculation_error:Free Cash Flow" in result.reasons
    assert result.score < 1.0


def test_formula_that_returns_text_loses_credit(workbook_case):
    workbook, _, addresses = workbook_case
    workbook.active[addresses["free_cash_flow"]] = '="not a number"'
    result = score_case(workbook_case)
    assert "formula_calculation_error:Free Cash Flow" in result.reasons
    assert result.score < 1.0


def test_unbounded_function_is_not_executed(workbook_case):
    workbook, _, addresses = workbook_case
    workbook.active[addresses["free_cash_flow"]] = '=LEN(REPT("x",1000000000))'
    with pytest.raises(FormulaEvaluationError, match="Unsupported spreadsheet function"):
        score_case(workbook_case)


def test_unbounded_nested_rounding_is_not_executed(workbook_case):
    workbook, _, addresses = workbook_case
    workbook.active[addresses["free_cash_flow"]] = (
        f"={addresses['cash_flow']}-{addresses['capex']}+0*ROUND(1,1000000000)"
    )
    with pytest.raises(FormulaEvaluationError, match="rounding precision"):
        score_case(workbook_case)


def test_excessive_worksheet_dimensions_are_unscored(workbook_case):
    workbook, _, _ = workbook_case
    workbook.active["XFD1048576"] = "distant cell"
    with pytest.raises(FormulaEvaluationError, match="size limit"):
        score_case(workbook_case)


def test_excessive_archive_members_are_unscored():
    buffer = BytesIO()
    with ZipFile(buffer, "w") as archive:
        for index in range(2001):
            archive.writestr(f"member-{index}", "")
    with pytest.raises(FormulaEvaluationError, match="archive exceeds"):
        score_workbook_bytes(buffer.getvalue(), {})
