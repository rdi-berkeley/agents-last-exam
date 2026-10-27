import asyncio
import json
import sys
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock

import pytest
from openpyxl import Workbook

from ale_run.tasks.driver import TaskDriver
from tasks.business_finance.equity_research_summary import main as task
from tasks.business_finance.equity_research_summary.scripts import score_workbook


@pytest.fixture
def episode():
    config = SimpleNamespace(
        metadata={
            "variant_name": "tsla_fy2023",
            "output_workbook": r"C:\output\TSLA_Financial_Summary.xlsx",
            "reference_manifest": r"C:\reference\reference_manifest.json",
        }
    )
    workbook = Workbook()
    workbook.active.title = "TSLA Summary"
    buffer = BytesIO()
    workbook.save(buffer)
    session = SimpleNamespace(
        file_exists=AsyncMock(return_value=True),
        directory_exists=AsyncMock(return_value=False),
        read_bytes=AsyncMock(return_value=buffer.getvalue()),
        read_file=AsyncMock(return_value='{"sheet_name_contains": "TSLA"}'),
    )
    return config, session


async def test_missing_reference_is_an_error_even_without_agent_output(episode):
    config, session = episode
    session.file_exists.return_value = False

    with pytest.raises(FileNotFoundError, match="Missing reference manifest"):
        await task.evaluate(config, session)

    session.read_bytes.assert_not_awaited()
    session.read_file.assert_not_awaited()


async def test_missing_agent_output_remains_zero(episode):
    config, session = episode
    session.file_exists.side_effect = lambda path: path == config.metadata["reference_manifest"]

    assert await task.evaluate(config, session) == [0.0]
    session.read_bytes.assert_not_awaited()
    session.read_file.assert_not_awaited()


@pytest.mark.parametrize("method", ["read_file", "read_bytes", "file_exists", "directory_exists"])
async def test_session_failures_propagate(episode, method):
    config, session = episode
    if method == "directory_exists":
        session.file_exists.return_value = False
    failure = OSError("storage unavailable")
    getattr(session, method).side_effect = failure

    with pytest.raises(OSError, match="storage unavailable") as caught:
        await task.evaluate(config, session)

    assert caught.value is failure


@pytest.mark.parametrize("manifest", ["", "{broken", "null", "[]"])
async def test_invalid_reference_is_not_agent_zero(episode, manifest):
    config, session = episode
    session.read_file.return_value = manifest

    with pytest.raises(ValueError):
        await task.evaluate(config, session)


@pytest.mark.parametrize("failure", [RuntimeError("scorer unavailable"), MemoryError("scorer OOM")])
async def test_scorer_failures_use_native_unscored_error(episode, monkeypatch, failure):
    config, session = episode
    scorer = AsyncMock(side_effect=failure)
    monkeypatch.setattr(task, "_score_workbook", scorer)

    with pytest.raises(type(failure)) as caught:
        await task.evaluate(config, session)
    assert caught.value is failure

    driver = object.__new__(TaskDriver)
    driver._session = session
    driver._task_loader = SimpleNamespace(get_evaluate_fn=lambda: task.evaluate)
    driver._make_task_cfg = lambda: config

    assert await driver.evaluate() == {
        "error": str(failure),
        "score": None,
        "success": None,
    }


@pytest.mark.parametrize("score", [0.0, 0.5, 1.0])
async def test_valid_scorer_results_are_unchanged(episode, monkeypatch, score):
    config, session = episode
    scorer = AsyncMock(return_value={"score": score, "passed": score == 1.0, "reasons": []})
    monkeypatch.setattr(task, "_score_workbook", scorer)

    assert await task.evaluate(config, session) == [score]
    scorer.assert_awaited_once_with(
        agent_bytes=session.read_bytes.return_value,
        manifest={"sheet_name_contains": "TSLA"},
    )


async def test_actual_workbook_still_scores(episode):
    assert await task.evaluate(*episode) == [1.0]


async def test_cancelled_evaluation_stops_the_scorer_process(monkeypatch):
    original_create = asyncio.create_subprocess_exec
    started = asyncio.Event()
    children = []

    async def start_slow_scorer(*args, **kwargs):
        process = await original_create(
            sys.executable, "-c", "import sys, time; sys.stdin.read(); time.sleep(60)", **kwargs
        )
        children.append(process)
        started.set()
        return process

    monkeypatch.setattr(task.asyncio, "create_subprocess_exec", start_slow_scorer)
    scoring = asyncio.create_task(task._score_workbook(b"workbook", {}))
    await started.wait()
    scoring.cancel()
    with pytest.raises(asyncio.CancelledError):
        await scoring
    assert children[0].returncode is not None


async def test_corrupt_agent_workbook_remains_zero(episode):
    config, session = episode
    session.read_bytes.return_value = b"not an xlsx workbook"

    assert await task.evaluate(config, session) == [0.0]


async def test_unassessable_formula_becomes_driver_error_not_zero(episode):
    config, session = episode
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "TSLA Summary"
    sheet["A1"], sheet["B1"] = "Operating Cash Flow", 40
    sheet["A2"], sheet["B2"] = "Capital Expenditures", 10
    sheet["A3"], sheet["B3"] = "Free Cash Flow", "=INDIRECT(X1)-B2"
    sheet["X1"] = "B1"
    buffer = BytesIO()
    workbook.save(buffer)
    session.read_bytes.return_value = buffer.getvalue()
    session.read_file.return_value = json.dumps(
        {
            "sheet_name_contains": "TSLA",
            "fixed_values": [
                {"label": "Operating Cash Flow", "value": 40},
                {"label": "Capital Expenditures", "value": 10},
            ],
            "formula_cells": [
                {
                    "label": "Free Cash Flow",
                    "must_reference": ["Operating Cash Flow", "Capital Expenditures"],
                }
            ],
        }
    )
    driver = object.__new__(TaskDriver)
    driver._session = session
    driver._task_loader = SimpleNamespace(get_evaluate_fn=lambda: task.evaluate)
    driver._make_task_cfg = lambda: config

    result = await driver.evaluate()
    assert result["score"] is None
    assert result["success"] is None
    assert result["error"]


@pytest.mark.parametrize("error", [MemoryError("no memory"), ValueError("dependency failure")])
def test_workbook_loader_infrastructure_error_is_not_zero(monkeypatch, error):
    monkeypatch.setattr(score_workbook, "load_workbook", Mock(side_effect=error))
    buffer = BytesIO()
    Workbook().save(buffer)

    with pytest.raises(type(error)):
        score_workbook.score_workbook_bytes(buffer.getvalue(), {})
